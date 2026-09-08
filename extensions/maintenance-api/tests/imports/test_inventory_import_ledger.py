from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from app.core.exceptions import BusinessValidationError, ConflictError
from app.db.base import Base
from app.models import (
    InventoryBalance,
    InventoryLedgerEntry,
    InventoryPolicy,
    InventoryTargetReceipt,
    InventoryTransaction,
    SparePart,
    Warehouse,
    WarehouseLocation,
)
from app.repositories.inventory_transaction_repository import (
    InventoryTransactionRepository,
)
from app.scripts.seed_master_data import seed as seed_master_data
from app.security.actor import MaintenanceRole
from app.services.import_service import master_data_import_service
from openpyxl import load_workbook
from sqlalchemy import func, select


def _append_named_row(sheet, values: dict[str, object]) -> None:
    headers = [cell.value for cell in sheet[1]]
    sheet.append([values.get(header) for header in headers])


def _inventory_workbook(
    *,
    operation: str = "CREATE",
    warehouse_code: str = "WH-IMPORT",
    spare_part_code: str = "SP-IMPORT",
    on_hand: str = "12.5000",
    reserved: str = "2.0000",
    damaged: str = "1.0000",
    quarantined: str = "0.5000",
    in_transit: str = "3.0000",
    safety_stock: str = "4.0000",
    reorder_point: str = "6.0000",
    maximum_stock: str = "20.0000",
    include_references: bool = False,
) -> bytes:
    workbook = load_workbook(
        BytesIO(master_data_import_service.template_bytes())
    )
    if include_references:
        _append_named_row(
            workbook["04_维修器材"],
            {
                "操作": "CREATE",
                "器材编码": spare_part_code,
                "器材名称": "Imported spare",
                "单位": "件",
            },
        )
        _append_named_row(
            workbook["07_库房"],
            {
                "操作": "CREATE",
                "库房编码": warehouse_code,
                "库房名称": "Imported warehouse",
                "库房状态": "NORMAL",
                "是否启用": True,
            },
        )
    _append_named_row(
        workbook["08_库存"],
        {
            "操作": operation,
            "库房编码": warehouse_code,
            "器材编码": spare_part_code,
            "现存数量": on_hand,
            "预留数量": reserved,
            "损坏数量": damaged,
            "隔离数量": quarantined,
            "在途数量": in_transit,
            "安全库存": safety_stock,
            "补货点": reorder_point,
            "最大库存": maximum_stock,
            "备注": "stable normalized import row",
        },
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _catalog(session, tenant_id: str = "tenant-a") -> tuple[Warehouse, SparePart]:
    warehouse = Warehouse(
        tenant_id=tenant_id,
        code="WH-IMPORT",
        name=f"{tenant_id} warehouse",
    )
    spare = SparePart(
        tenant_id=tenant_id,
        code="SP-IMPORT",
        name=f"{tenant_id} spare",
        unit="件",
    )
    session.add_all([warehouse, spare])
    session.flush()
    return warehouse, spare


def _apply(session, actor_admin, *, task_id: str, content: bytes):
    return master_data_import_service.apply(
        session,
        actor=actor_admin,
        task_id=task_id,
        content=content,
        filename="inventory.xlsx",
    )


def _sync_apply(session, actor_admin, *, content: bytes):
    return master_data_import_service.apply(
        session,
        actor=actor_admin,
        content=content,
        filename="inventory.xlsx",
    )


def _xlsx_byte_variant(content: bytes, marker: str) -> bytes:
    stream = BytesIO(content)
    with ZipFile(stream, "a", compression=ZIP_DEFLATED) as archive:
        archive.writestr(f"custom/retry-{marker}.txt", marker)
    return stream.getvalue()


def test_inventory_create_writes_policy_default_opening_and_five_component_ledger(
    session,
    actor_admin,
):
    warehouse, spare = _catalog(session)

    result = _apply(
        session,
        actor_admin,
        task_id="task-opening",
        content=_inventory_workbook(),
    )
    session.commit()

    policy = session.scalar(select(InventoryPolicy))
    location = session.scalar(select(WarehouseLocation))
    balance = session.scalar(select(InventoryBalance))
    transaction = session.scalar(select(InventoryTransaction))
    entry = session.scalar(select(InventoryLedgerEntry))
    assert result.created["08_库存"] == 1
    assert policy is not None
    assert (policy.warehouse_id, policy.spare_part_id) == (
        warehouse.id,
        spare.id,
    )
    assert location is not None and location.code == "DEFAULT"
    assert balance is not None
    assert (
        balance.on_hand_quantity,
        balance.reserved_quantity,
        balance.damaged_quantity,
        balance.quarantined_quantity,
        balance.in_transit_quantity,
    ) == tuple(
        Decimal(value)
        for value in ("12.5000", "2", "1", "0.5000", "3")
    )
    assert transaction is not None
    assert transaction.operation_type == "OPENING"
    assert transaction.idempotency_key == "import:task-opening:08_库存:2"
    assert entry is not None
    assert (
        entry.on_hand_delta,
        entry.reserved_delta,
        entry.damaged_delta,
        entry.quarantined_delta,
        entry.in_transit_delta,
    ) == tuple(
        Decimal(value)
        for value in ("12.5000", "2", "1", "0.5000", "3")
    )


def test_inventory_update_adjusts_existing_aggregate_to_target(
    session,
    actor_admin,
):
    warehouse, spare = _catalog(session)
    default = WarehouseLocation(
        tenant_id=actor_admin.tenant_id,
        warehouse_id=warehouse.id,
        code="DEFAULT",
        name="Default location",
        location_type="DEFAULT",
    )
    shelf = WarehouseLocation(
        tenant_id=actor_admin.tenant_id,
        warehouse_id=warehouse.id,
        code="SHELF-A",
        name="Shelf A",
        location_type="STORAGE",
    )
    session.add_all([default, shelf])
    session.flush()
    session.add(
        InventoryPolicy(
            tenant_id=actor_admin.tenant_id,
            warehouse_id=warehouse.id,
            spare_part_id=spare.id,
            safety_stock=1,
            reorder_point=2,
            maximum_stock=30,
        )
    )
    session.add_all(
        [
            InventoryBalance(
                tenant_id=actor_admin.tenant_id,
                warehouse_id=warehouse.id,
                location_id=default.id,
                spare_part_id=spare.id,
                on_hand_quantity=5,
                reserved_quantity=1,
                in_transit_quantity=1,
            ),
            InventoryBalance(
                tenant_id=actor_admin.tenant_id,
                warehouse_id=warehouse.id,
                location_id=shelf.id,
                spare_part_id=spare.id,
                on_hand_quantity=2,
                damaged_quantity=1,
            ),
        ]
    )
    session.commit()

    _apply(
        session,
        actor_admin,
        task_id="task-adjust",
        content=_inventory_workbook(
            operation="UPDATE",
            on_hand="10",
            reserved="2",
            damaged="1",
            quarantined="1",
            in_transit="4",
        ),
    )
    session.commit()

    balances = session.scalars(
        select(InventoryBalance).order_by(InventoryBalance.id)
    ).all()
    assert tuple(
        sum(getattr(item, field) for item in balances)
        for field in (
            "on_hand_quantity",
            "reserved_quantity",
            "damaged_quantity",
            "quarantined_quantity",
            "in_transit_quantity",
        )
    ) == tuple(Decimal(value) for value in ("10", "2", "1", "1", "4"))
    transaction = session.scalar(select(InventoryTransaction))
    assert transaction is not None
    assert transaction.operation_type == "ADJUST"
    assert transaction.idempotency_key == "import:task-adjust:08_库存:2"


def test_same_task_replay_is_zero_duplicate_and_changed_payload_conflicts(
    session,
    actor_admin,
):
    _catalog(session)
    original = _inventory_workbook()
    _apply(session, actor_admin, task_id="task-replay", content=original)
    session.commit()

    _apply(session, actor_admin, task_id="task-replay", content=original)
    session.commit()
    assert session.scalar(select(func.count()).select_from(InventoryTransaction)) == 1
    assert session.scalar(select(func.count()).select_from(InventoryLedgerEntry)) == 1

    with pytest.raises(ConflictError) as raised:
        _apply(
            session,
            actor_admin,
            task_id="task-replay",
            content=_inventory_workbook(on_hand="13.5000"),
        )
    assert raised.value.code == "IDEMPOTENCY_KEY_REUSED"


def test_sync_logically_identical_byte_different_workbook_replays_same_receipt(
    session,
    actor_admin,
):
    _catalog(session)
    original = _inventory_workbook()
    byte_variant = _xlsx_byte_variant(original, "different-container-bytes")
    assert original != byte_variant

    _sync_apply(session, actor_admin, content=original)
    session.commit()
    _sync_apply(session, actor_admin, content=byte_variant)
    session.commit()

    assert session.scalar(select(func.count()).select_from(InventoryTargetReceipt)) == 1
    assert session.scalar(select(func.count()).select_from(InventoryTransaction)) == 1
    assert session.scalar(select(func.count()).select_from(InventoryLedgerEntry)) == 1


def test_sync_command_identity_is_canonical_but_row_order_sensitive():
    assert hasattr(master_data_import_service, "_synchronous_task_id")
    first = {
        "08_库存": [
            {"_row": 2, "operation": "UPDATE", "on_hand_quantity": Decimal("1")},
            {"_row": 3, "operation": "UPDATE", "on_hand_quantity": Decimal("2")},
        ]
    }
    logically_same = {
        "08_库存": [
            {"on_hand_quantity": Decimal("1.0000"), "operation": "UPDATE", "_row": 2},
            {"on_hand_quantity": Decimal("2.0000"), "operation": "UPDATE", "_row": 3},
        ]
    }
    changed_cell = {
        "08_库存": [
            {"_row": 2, "operation": "UPDATE", "on_hand_quantity": Decimal("9")},
            {"_row": 3, "operation": "UPDATE", "on_hand_quantity": Decimal("2")},
        ]
    }
    reordered = {"08_库存": list(reversed(first["08_库存"]))}
    identity = master_data_import_service._synchronous_task_id(
        normalized=first,
        mapping={"08_库存": {"现存数量": "on_hand_quantity"}},
        template_version="1.0",
    )
    assert identity == master_data_import_service._synchronous_task_id(
        normalized=logically_same,
        mapping={"08_库存": {"现存数量": "on_hand_quantity"}},
        template_version="1.0",
    )
    assert identity != master_data_import_service._synchronous_task_id(
        normalized=changed_cell,
        mapping={"08_库存": {"现存数量": "on_hand_quantity"}},
        template_version="1.0",
    )
    assert identity != master_data_import_service._synchronous_task_id(
        normalized=reordered,
        mapping={"08_库存": {"现存数量": "on_hand_quantity"}},
        template_version="1.0",
    )
    assert identity != master_data_import_service._synchronous_task_id(
        normalized=first,
        mapping={"08_库存": {"现存数量": "on_hand_quantity"}},
        template_version="2.0",
    )


def test_explicit_queued_task_uuid_remains_receipt_identity_across_zip_bytes(
    session,
    actor_admin,
):
    _catalog(session)
    original = _inventory_workbook()
    _apply(session, actor_admin, task_id="immutable-task-uuid", content=original)
    session.commit()
    _apply(
        session,
        actor_admin,
        task_id="immutable-task-uuid",
        content=_xlsx_byte_variant(original, "queued-retry"),
    )
    session.commit()

    receipt = session.scalar(select(InventoryTargetReceipt))
    assert receipt is not None
    assert receipt.idempotency_key == "import:immutable-task-uuid:08_库存:2"
    assert session.scalar(select(func.count()).select_from(InventoryTargetReceipt)) == 1


def test_zero_quantity_row_still_records_source_receipt_and_rejects_reuse(
    session,
    actor_admin,
):
    _catalog(session)
    zero = _inventory_workbook(
        on_hand="0",
        reserved="0",
        damaged="0",
        quarantined="0",
        in_transit="0",
    )
    _apply(session, actor_admin, task_id="task-zero", content=zero)
    session.commit()

    receipt_table = Base.metadata.tables.get("inventory_target_receipts")
    assert receipt_table is not None
    receipt = session.execute(select(receipt_table)).mappings().one()
    assert receipt["idempotency_key"] == "import:task-zero:08_库存:2"
    assert receipt["status"] == "COMPLETED"
    assert receipt["result_json"] == {
        "created_identity": True,
        "operation_type": None,
        "transaction_id": None,
    }
    assert session.scalar(select(func.count()).select_from(InventoryTransaction)) == 0
    assert session.scalar(select(func.count()).select_from(InventoryLedgerEntry)) == 0

    with pytest.raises(ConflictError) as raised:
        _apply(
            session,
            actor_admin,
            task_id="task-zero",
            content=_inventory_workbook(
                on_hand="0",
                reserved="0",
                damaged="0",
                quarantined="0",
                in_transit="0",
                safety_stock="5",
                reorder_point="6",
            ),
        )
    assert raised.value.code == "IDEMPOTENCY_KEY_REUSED"


def test_inventory_import_is_tenant_scoped_even_with_shared_codes(
    session,
    actor_context,
):
    _catalog(session, "tenant-a")
    warehouse_b, spare_b = _catalog(session, "tenant-b")
    location_b = WarehouseLocation(
        tenant_id="tenant-b",
        warehouse_id=warehouse_b.id,
        code="DEFAULT",
        name="Default",
        location_type="DEFAULT",
    )
    session.add(location_b)
    session.flush()
    session.add_all(
        [
            InventoryPolicy(
                tenant_id="tenant-b",
                warehouse_id=warehouse_b.id,
                spare_part_id=spare_b.id,
            ),
            InventoryBalance(
                tenant_id="tenant-b",
                warehouse_id=warehouse_b.id,
                location_id=location_b.id,
                spare_part_id=spare_b.id,
                on_hand_quantity=99,
            ),
        ]
    )
    session.commit()

    actor_a = actor_context(
        tenant_id="tenant-a",
        role=MaintenanceRole.ADMIN,
    )
    _apply(session, actor_a, task_id="task-tenant", content=_inventory_workbook())
    session.commit()

    tenant_b_balance = session.scalar(
        select(InventoryBalance).where(InventoryBalance.tenant_id == "tenant-b")
    )
    assert tenant_b_balance is not None
    assert tenant_b_balance.on_hand_quantity == 99
    assert session.scalar(
        select(func.count())
        .select_from(InventoryTransaction)
        .where(InventoryTransaction.tenant_id == "tenant-b")
    ) == 0


def test_inventory_import_failure_rolls_back_identity_policy_and_ledger(
    session,
    actor_admin,
    monkeypatch,
):
    _catalog(session)

    def fail_append(*args, **kwargs):
        del args, kwargs
        raise RuntimeError("ledger append failed")

    monkeypatch.setattr(
        InventoryTransactionRepository,
        "append_entry",
        fail_append,
    )
    with pytest.raises(RuntimeError, match="ledger append failed"):
        _apply(
            session,
            actor_admin,
            task_id="task-rollback",
            content=_inventory_workbook(),
        )
    session.rollback()

    for model in (
        InventoryPolicy,
        WarehouseLocation,
        InventoryBalance,
        InventoryTransaction,
        InventoryLedgerEntry,
        InventoryTargetReceipt,
    ):
        assert session.scalar(select(func.count()).select_from(model)) == 0


@pytest.mark.parametrize(
    ("field", "target"),
    [
        ("on_hand", "9"),
        ("reserved", "1"),
        ("damaged", "1"),
        ("quarantined", "1"),
        ("in_transit", "1"),
    ],
)
def test_inventory_target_below_non_default_floor_is_rejected_before_mutation(
    session,
    actor_admin,
    field,
    target,
):
    warehouse, spare = _catalog(session)
    default = WarehouseLocation(
        tenant_id=actor_admin.tenant_id,
        warehouse_id=warehouse.id,
        code="DEFAULT",
        name="Default location",
        location_type="DEFAULT",
    )
    shelf = WarehouseLocation(
        tenant_id=actor_admin.tenant_id,
        warehouse_id=warehouse.id,
        code="SHELF",
        name="Shelf",
        location_type="STORAGE",
    )
    session.add_all([default, shelf])
    session.flush()
    policy = InventoryPolicy(
        tenant_id=actor_admin.tenant_id,
        warehouse_id=warehouse.id,
        spare_part_id=spare.id,
        safety_stock=Decimal("1"),
        reorder_point=Decimal("2"),
        maximum_stock=Decimal("30"),
        notes="unchanged",
    )
    session.add_all(
        [
            policy,
            InventoryBalance(
                tenant_id=actor_admin.tenant_id,
                warehouse_id=warehouse.id,
                location_id=default.id,
                spare_part_id=spare.id,
                on_hand_quantity=Decimal("3"),
            ),
            InventoryBalance(
                tenant_id=actor_admin.tenant_id,
                warehouse_id=warehouse.id,
                location_id=shelf.id,
                spare_part_id=spare.id,
                on_hand_quantity=Decimal("10"),
                reserved_quantity=Decimal("2"),
                damaged_quantity=Decimal("2"),
                quarantined_quantity=Decimal("2"),
                in_transit_quantity=Decimal("2"),
            ),
        ]
    )
    session.commit()
    targets = {
        "on_hand": "13",
        "reserved": "2",
        "damaged": "2",
        "quarantined": "2",
        "in_transit": "2",
    }
    targets[field] = target

    with pytest.raises(BusinessValidationError) as raised:
        _apply(
            session,
            actor_admin,
            task_id=f"task-unreachable-{field}",
            content=_inventory_workbook(
                operation="UPDATE",
                on_hand=targets["on_hand"],
                reserved=targets["reserved"],
                damaged=targets["damaged"],
                quarantined=targets["quarantined"],
                in_transit=targets["in_transit"],
                safety_stock="4",
                reorder_point="6",
                maximum_stock="40",
            ),
        )
    assert raised.value.code == "INVENTORY_TARGET_UNREACHABLE"
    assert raised.value.details == {
        "unreachable_components": [field],
        "retryable": False,
    }
    session.rollback()
    session.refresh(policy)
    assert (policy.safety_stock, policy.reorder_point, policy.maximum_stock, policy.notes) == (
        Decimal("1"),
        Decimal("2"),
        Decimal("30"),
        "unchanged",
    )
    assert session.scalar(select(func.count()).select_from(InventoryTargetReceipt)) == 0
    assert session.scalar(select(func.count()).select_from(InventoryTransaction)) == 0
    assert session.scalar(select(func.count()).select_from(InventoryLedgerEntry)) == 0


def test_inventory_sync_execute_requires_admin_and_admin_writes_ledger(
    client,
    internal_auth_headers,
    session,
):
    content = _inventory_workbook(include_references=True)
    path = "/api/v1/master-data/import/execute"

    denied = client.post(
        path,
        files={"file": ("inventory.xlsx", content)},
        headers=internal_auth_headers(role=MaintenanceRole.CONTRIBUTOR),
    )
    assert denied.status_code == 403
    for model in (
        InventoryPolicy,
        InventoryBalance,
        InventoryTransaction,
        InventoryLedgerEntry,
    ):
        assert session.scalar(select(func.count()).select_from(model)) == 0

    allowed = client.post(
        path,
        files={"file": ("inventory.xlsx", content)},
        headers=internal_auth_headers(role=MaintenanceRole.ADMIN),
    )
    assert allowed.status_code == 200
    assert session.scalar(select(func.count()).select_from(InventoryLedgerEntry)) == 1


def test_seed_inventory_uses_fixed_receipts_and_reruns_without_duplicates(session):
    seed_master_data(tenant_id="tenant-seed")
    first_transactions = session.scalar(
        select(func.count())
        .select_from(InventoryTransaction)
        .where(InventoryTransaction.tenant_id == "tenant-seed")
    )
    first_entries = session.scalar(
        select(func.count())
        .select_from(InventoryLedgerEntry)
        .where(InventoryLedgerEntry.tenant_id == "tenant-seed")
    )
    keys = set(
        session.scalars(
            select(InventoryTransaction.idempotency_key).where(
                InventoryTransaction.tenant_id == "tenant-seed"
            )
        )
    )
    warehouse_ids = set(
        session.scalars(
            select(Warehouse.id).where(Warehouse.tenant_id == "tenant-seed")
        )
    )
    spare_ids = set(
        session.scalars(
            select(SparePart.id).where(
                SparePart.tenant_id == "tenant-seed",
                SparePart.code.in_([f"SP-{index:03d}" for index in range(1, 11)]),
            )
        )
    )
    assert first_transactions == first_entries == 30
    assert keys == {
        f"seed:inventory:tenant-seed:{warehouse_id}:{spare_id}"
        for warehouse_id in warehouse_ids
        for spare_id in spare_ids
    }

    seed_master_data(tenant_id="tenant-seed")
    session.expire_all()
    assert session.scalar(
        select(func.count())
        .select_from(InventoryTransaction)
        .where(InventoryTransaction.tenant_id == "tenant-seed")
    ) == first_transactions
    assert session.scalar(
        select(func.count())
        .select_from(InventoryLedgerEntry)
        .where(InventoryLedgerEntry.tenant_id == "tenant-seed")
    ) == first_entries


def test_runtime_mapper_registry_has_no_removed_inventory_table():
    assert "warehouse_inventories" not in Base.metadata.tables
