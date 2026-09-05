from __future__ import annotations

from collections.abc import Callable
from io import BytesIO
from types import SimpleNamespace

from app.models import (
    AIReportCitation,
    AIReportJob,
    AIReportSection,
    AIReportVersion,
)
from app.models.enums import (
    AIReportJobStatus,
    AIReportType,
    AIReportVersionStatus,
)
from app.security.actor import MaintenanceRole
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

_COMPACT_JWT = (
    "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9."
    "eyJzdWIiOiIxMjM0NTY3ODkwIiwibmFtZSI6IkpvaG4gRG9lIn0."
    "SflKxwRJSMeKKF2QT4fwpMeJf36POk6yJV_adQssw5c"
)


def _viewer_headers(
    internal_auth_headers: Callable[..., dict[str, str]],
) -> dict[str, str]:
    return internal_auth_headers(
        tenant_id="tenant-safe-provenance",
        user_id="safe-provenance-viewer",
        role=MaintenanceRole.VIEWER,
    )


def _document_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    text = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            text.extend(cell.text for cell in row.cells)
    return "\n".join(text)


def test_report_detail_exposes_only_public_source_projection(
    client: TestClient,
    session: Session,
    internal_auth_headers: Callable[..., dict[str, str]],
    monkeypatch,
    tmp_path,
) -> None:
    job = AIReportJob(
        tenant_id="tenant-safe-provenance",
        report_code="AIR-SAFE-PROVENANCE",
        report_type=AIReportType.MANAGEMENT_DECISION,
        status=AIReportJobStatus.CREATED,
        title="Safe provenance",
        progress_percent=0,
    )
    session.add(job)
    session.flush()
    version = AIReportVersion(
        tenant_id=job.tenant_id,
        report_job_id=job.id,
        version_number=1,
        status=AIReportVersionStatus.DRAFT,
        template_version="1.0",
        content_digest="a" * 64,
        metadata_json={
            "purpose": "Safe purpose",
            "tenant_id": "metadata-tenant-must-not-leak",
            "internal_metadata": {
                "provider_token": "metadata-token-must-not-leak"
            },
            "source_path": "C:/metadata-path-must-not-leak",
            "display": {
                "label": _COMPACT_JWT,
                "location_note": "Stored at /report.json (archived)",
                "safe_label": "Safe nested label",
                "provider_token": "nested-token-must-not-leak",
                "source_snapshot_json": {
                    "ordinary_value": "nested snapshot payload must not leak"
                },
                "database_record_json": {
                    "ordinary_value": "nested record payload must not leak"
                },
                "safe_sibling": "Safe metadata sibling",
            },
            "safe_tags": ["Safe tag", _COMPACT_JWT],
            "notes": [
                "Before\n/report.json\nAfter",
                "Safe root-level note",
            ],
            "safe_nested_list": [
                "Safe list sibling",
                {
                    "source_snapshot_json": {
                        "ordinary_value": "list snapshot payload must not leak"
                    },
                    "database_record_json": {
                        "ordinary_value": "list record payload must not leak"
                    },
                    "safe_sibling": "Safe list mapping sibling",
                },
            ],
            "_section_tables": {
                "summary": [
                    {
                        "title": _COMPACT_JWT,
                        "columns": [
                            "Part",
                            "tenant-id-column-must-not-leak",
                            _COMPACT_JWT,
                            "Quantity",
                        ],
                        "rows": [
                            [
                                "Widget",
                                "provider-credential-cell-must-not-leak",
                                _COMPACT_JWT,
                                3,
                            ]
                        ],
                        "tenant_id": "table-tenant-must-not-leak",
                        "internal_metadata": {
                            "provider_token": "table-token-must-not-leak"
                        },
                        "evidence": {
                            "credential": "table-credential-must-not-leak"
                        },
                        "source_snapshot_json": {
                            "tenant_id": "table-snapshot-must-not-leak"
                        },
                        "database_record_json": {
                            "token": "table-record-must-not-leak"
                        },
                    }
                ]
            },
            "_section_citations": {
                "summary": [
                    "E-SAFE",
                    "tenant-id-section-citation-must-not-leak",
                    _COMPACT_JWT,
                    "provider-credential-section-citation-must-not-leak",
                    "secret-section-citation-must-not-leak",
                    {
                        "provider_token": "section-citation-token-must-not-leak"
                    },
                ]
            },
        },
        source_snapshot_json={
            "schema_version": "1.1",
            "capture_mode": "AUTHORITATIVE_CREATE",
            "provenance_completeness": "AUTHORITATIVE",
            "tenant_id": "tenant-id-must-not-leak",
            "provider_token": "provider-token-must-not-leak",
            "sources": [
                {
                    "type": "AI_SESSION",
                    "id": "42",
                    "version": "7",
                    "lineage_id": "session-42",
                    "digest": "b" * 64,
                    "evidence": {
                        "credential": "credential-must-not-leak"
                    },
                }
            ],
        },
    )
    session.add(version)
    session.flush()
    session.add(
        AIReportSection(
            tenant_id=job.tenant_id,
            report_version_id=version.id,
            section_code="summary",
            title="Safe section",
            order_index=0,
            content="Before (/report.json) after",
            source_type="DETERMINISTIC",
        )
    )
    session.add(
        AIReportCitation(
            tenant_id=job.tenant_id,
            report_version_id=version.id,
            citation_id="E-SAFE",
            source_type="WEKNORA_DOCUMENT",
            source_name="Safe evidence",
            page_number=4,
            database_record_json={
                "evidence": {
                    "provider_token": "citation-token-must-not-leak"
                },
                "tenant_id": "citation-tenant-must-not-leak",
            },
        )
    )
    session.add(
        AIReportCitation(
            tenant_id=job.tenant_id,
            report_version_id=version.id,
            citation_id="E-ROOT-PATH",
            source_type="WEKNORA_DOCUMENT",
            source_name="Before\n/report.json\nAfter",
        )
    )
    session.add(
        AIReportCitation(
            tenant_id=job.tenant_id,
            report_version_id=version.id,
            citation_id="E-REDACTED",
            source_type="WEKNORA_DOCUMENT",
            source_name=_COMPACT_JWT,
            document_version="2026.09",
            page_number=7,
            chunk_reference="C:/citation-path-must-not-leak",
            knowledge_node="provider-token-citation-must-not-leak",
            database_record_json={
                "provider_token": "citation-record-token-must-not-leak"
            },
        )
    )
    session.commit()

    response = client.get(
        f"/api/v1/reports/{job.id}",
        headers=_viewer_headers(internal_auth_headers),
    )

    assert response.status_code == 200
    body = response.text
    for unsafe in (
        "source_snapshot_json",
        "tenant-id-must-not-leak",
        "provider-token-must-not-leak",
        "credential-must-not-leak",
        "metadata-tenant-must-not-leak",
        "metadata-token-must-not-leak",
        "metadata-path-must-not-leak",
        "nested-token-must-not-leak",
        "Stored at /report.json (archived)",
        "Before (/report.json) after",
        "Before\n/report.json\nAfter",
        "nested snapshot payload must not leak",
        "nested record payload must not leak",
        "list snapshot payload must not leak",
        "list record payload must not leak",
        "source_snapshot_json",
        "database_record_json",
        "Stored at C:/reports/report.json (archived)",
        "Stored at /var/lib/reports/report.json (archived)",
        "Stored at \\\\server\\share\\report.json (archived)",
        "Stored at file:///var/lib/reports/report.json (archived)",
        "Before (C:/reports/report.json) after",
        "citation-token-must-not-leak",
        "citation-tenant-must-not-leak",
        "citation-path-must-not-leak",
        "provider-token-citation-must-not-leak",
        "citation-record-token-must-not-leak",
        "database_record_json",
        "table-tenant-must-not-leak",
        "table-token-must-not-leak",
        "table-credential-must-not-leak",
        "table-snapshot-must-not-leak",
        "table-record-must-not-leak",
        "section-citation-token-must-not-leak",
        "tenant-id-column-must-not-leak",
        "jwt-column-must-not-leak",
        "provider-credential-cell-must-not-leak",
        "secret-cell-must-not-leak",
        "tenant-id-section-citation-must-not-leak",
        "jwt-section-citation-must-not-leak",
        "provider-credential-section-citation-must-not-leak",
        "secret-section-citation-must-not-leak",
        _COMPACT_JWT,
    ):
        assert unsafe not in body
    assert response.json()["data"]["source_versions"] == {
        "capture_mode": "AUTHORITATIVE_CREATE",
        "provenance_completeness": "AUTHORITATIVE",
        "sources": [
            {
                "type": "AI_SESSION",
                "id": "42",
                "version": "7",
                "lineage_id": "session-42",
                "digest": "b" * 64,
            }
        ],
    }
    detail = response.json()["data"]
    assert detail["metadata"] == {
        "purpose": "Safe purpose",
        "display": {
            "safe_label": "Safe nested label",
            "safe_sibling": "Safe metadata sibling",
        },
        "safe_tags": ["Safe tag"],
        "notes": ["Safe root-level note"],
        "safe_nested_list": [
            "Safe list sibling",
            {"safe_sibling": "Safe list mapping sibling"},
        ],
    }
    assert detail["citations"] == [
        {
            "citation_id": "E-REDACTED",
            "source_type": "WEKNORA_DOCUMENT",
            "source_name": None,
            "document_version": "2026.09",
            "page_number": 7,
            "chunk_reference": None,
            "knowledge_node": None,
        },
        {
            "citation_id": "E-ROOT-PATH",
            "source_type": "WEKNORA_DOCUMENT",
            "source_name": None,
            "document_version": None,
            "page_number": None,
            "chunk_reference": None,
            "knowledge_node": None,
        },
        {
            "citation_id": "E-SAFE",
            "source_type": "WEKNORA_DOCUMENT",
            "source_name": "Safe evidence",
            "document_version": None,
            "page_number": 4,
            "chunk_reference": None,
            "knowledge_node": None,
        },
    ]
    assert detail["sections"] == [
        {
            "section_code": "summary",
            "title": "Safe section",
            "content": None,
            "source_type": "DETERMINISTIC",
            "citations": ["E-SAFE"],
            "tables": [
                {
                    "columns": ["Part", "", "", "Quantity"],
                    "rows": [["Widget", "", "", 3]],
                }
            ],
        }
    ]
    monkeypatch.setattr(
        "app.services.ai_report_service.get_settings",
        lambda: SimpleNamespace(ai_report_export_dir=str(tmp_path)),
    )
    exports = {
        "json": client.get(
            f"/api/v1/reports/{job.id}/exports/json",
            headers=_viewer_headers(internal_auth_headers),
        ).text,
        "markdown": client.get(
            f"/api/v1/reports/{job.id}/exports/markdown",
            headers=_viewer_headers(internal_auth_headers),
        ).text,
        "docx": _document_text(
            client.get(
                f"/api/v1/reports/{job.id}/exports/docx",
                headers=_viewer_headers(internal_auth_headers),
            ).content
        ),
    }
    for output in exports.values():
        assert "Safe purpose" in output
        assert "Safe tag" in output
        assert "Safe root-level note" in output
        assert "Safe metadata sibling" in output
        assert "Safe list sibling" in output
        assert "Safe list mapping sibling" in output
        assert "Safe evidence" in output
        assert "Part" in output
        assert "Widget" in output
        assert "E-SAFE" in output
        for unsafe in (
            "metadata-tenant-must-not-leak",
            "metadata-token-must-not-leak",
            "metadata-path-must-not-leak",
            "nested-token-must-not-leak",
            "Stored at /report.json (archived)",
            "Before (/report.json) after",
            "Before\n/report.json\nAfter",
            "nested snapshot payload must not leak",
            "nested record payload must not leak",
            "list snapshot payload must not leak",
            "list record payload must not leak",
            "source_snapshot_json",
            "database_record_json",
            "Stored at C:/reports/report.json (archived)",
            "Stored at /var/lib/reports/report.json (archived)",
            "Stored at \\\\server\\share\\report.json (archived)",
            "Stored at file:///var/lib/reports/report.json (archived)",
            "Before (C:/reports/report.json) after",
            "citation-token-must-not-leak",
            "citation-tenant-must-not-leak",
            "citation-path-must-not-leak",
            "provider-token-citation-must-not-leak",
            "citation-record-token-must-not-leak",
            "database_record_json",
            "table-tenant-must-not-leak",
            "table-token-must-not-leak",
            "table-credential-must-not-leak",
            "table-snapshot-must-not-leak",
            "table-record-must-not-leak",
            "section-citation-token-must-not-leak",
            "tenant-id-column-must-not-leak",
            "jwt-column-must-not-leak",
            "provider-credential-cell-must-not-leak",
            "secret-cell-must-not-leak",
            "tenant-id-section-citation-must-not-leak",
            "jwt-section-citation-must-not-leak",
            "provider-credential-section-citation-must-not-leak",
            "secret-section-citation-must-not-leak",
            _COMPACT_JWT,
        ):
            assert unsafe not in output


def test_report_export_filenames_remain_versioned(
    client: TestClient,
    session: Session,
    internal_auth_headers: Callable[..., dict[str, str]],
    monkeypatch,
    tmp_path,
) -> None:
    job = AIReportJob(
        tenant_id="tenant-safe-provenance",
        report_code="AIR-SAFE-FILENAME",
        report_type=AIReportType.MANAGEMENT_DECISION,
        status=AIReportJobStatus.CREATED,
        title="Safe export filename",
        progress_percent=0,
    )
    session.add(job)
    session.flush()
    session.add(
        AIReportVersion(
            tenant_id=job.tenant_id,
            report_job_id=job.id,
            version_number=3,
            status=AIReportVersionStatus.DRAFT,
            template_version="1.0",
            content_digest="c" * 64,
        )
    )
    session.commit()
    monkeypatch.setattr(
        "app.services.ai_report_service.get_settings",
        lambda: SimpleNamespace(ai_report_export_dir=str(tmp_path)),
    )

    for export_format, extension in (
        ("json", "json"),
        ("markdown", "md"),
        ("docx", "docx"),
    ):
        response = client.get(
            f"/api/v1/reports/{job.id}/exports/{export_format}",
            headers=_viewer_headers(internal_auth_headers),
        )

        assert response.status_code == 200
        assert response.headers["content-disposition"] == (
            f'attachment; filename="AIR-SAFE-FILENAME-v3.{extension}"'
        )
        assert (tmp_path / f"AIR-SAFE-FILENAME-v3.{extension}").is_file()
