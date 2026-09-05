from __future__ import annotations

import json
from io import BytesIO

from app.exporters.ai_report_docx import export_report_docx
from app.exporters.ai_report_json import export_report_json
from app.exporters.ai_report_markdown import export_report_markdown
from docx import Document


def _document_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    text = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            text.extend(cell.text for cell in row.cells)
    return "\n".join(text)


def _report_payload() -> dict:
    source_versions = {
        "capture_mode": "AUTHORITATIVE_CREATE",
        "provenance_completeness": "AUTHORITATIVE",
        "sources": [
            {
                "type": "AI_SESSION",
                "id": "42",
                "version": "7",
                "lineage_id": "session-42",
                "digest": "b" * 64,
            }
        ],
    }
    return {
        "report_code": "AIR-SAFE-PROVENANCE",
        "report_type": "MANAGEMENT_DECISION",
        "title": "Safe provenance export",
        "version_number": 3,
        "template_version": "1.0",
        "generated_at": "2026-09-05T08:30:00+00:00",
        "generation_mode": "RULE_FALLBACK",
        "input_digest": "a" * 64,
        "source_versions": source_versions,
        "citations": [
            {
                "citation_id": "E-42",
                "source_name": "Safe evidence",
                "page_number": 4,
            }
        ],
        "metadata": {"purpose": "safe export"},
        "sections": [],
        "source_snapshot_json": {
            "tenant_id": "tenant-id-must-not-leak",
            "provider_token": "provider-token-must-not-leak",
            "credential": "credential-must-not-leak",
        },
        "metadata_json": {"internal": "internal-must-not-leak"},
        "file_path": "C:/absolute/path-must-not-leak",
    }


def test_exports_render_the_safe_serialized_provenance_only() -> None:
    report = _report_payload()
    json_export = json.loads(export_report_json(report))
    markdown_export = export_report_markdown(report)
    docx_export = _document_text(export_report_docx(report))

    assert json_export["provenance"] == {
        "report_code": report["report_code"],
        "report_type": report["report_type"],
        "version_number": report["version_number"],
        "template_version": report["template_version"],
        "generated_at": report["generated_at"],
        "generation_mode": report["generation_mode"],
        "input_digest": report["input_digest"],
        "source_versions": report["source_versions"],
        "citations": report["citations"],
    }
    for output in (json.dumps(json_export), markdown_export, docx_export):
        assert "AI_SESSION" in output
        assert "b" * 64 in output
        for unsafe in (
            "source_snapshot_json",
            "tenant-id-must-not-leak",
            "provider-token-must-not-leak",
            "credential-must-not-leak",
            "internal-must-not-leak",
            "C:/absolute/path-must-not-leak",
        ):
            assert unsafe not in output


def test_exports_render_unavailable_provenance_for_a_malformed_snapshot() -> None:
    report = _report_payload()
    report["source_versions"] = {}

    json_export = json.loads(export_report_json(report))
    markdown_export = export_report_markdown(report)
    docx_export = _document_text(export_report_docx(report))

    assert json_export["provenance"]["source_versions"] == {}
    assert "Source versions / hashes**: Unavailable" in markdown_export
    assert "Source versions / hashes" in docx_export
    assert "Unavailable" in docx_export
    for output in (json.dumps(json_export), markdown_export, docx_export):
        assert "provider-token-must-not-leak" not in output
